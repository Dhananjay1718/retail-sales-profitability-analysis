"""Reproducible synthetic retail analytics pipeline. Run from any directory."""
from pathlib import Path
import html
import json
import platform
import sqlite3
import zipfile

import numpy as np
import pandas as pd
import plotly
import xlsxwriter
import openpyxl
import nbformat
from nbclient import NotebookClient

ROOT = Path(__file__).resolve().parents[1]
for folder in ("data", "reports", "excel", "dashboard", "dist", "sql"):
    (ROOT / folder).mkdir(exist_ok=True)

def write(path, value):
    (ROOT / path).write_text(value.strip() + "\n", encoding="utf-8")

# Generate a stable, explicitly synthetic case. One row = one complete order.
rng = np.random.default_rng(42)
catalog = pd.DataFrame({
    "product": ["Laptop", "Headphones", "Keyboard", "Office Chair",
                "Desk", "Notebook", "Pen Set", "Backpack"],
    "category": ["Electronics", "Electronics", "Electronics", "Furniture",
                 "Furniture", "Stationery", "Stationery", "Accessories"],
    "unit_price": [50000, 2500, 1800, 6500, 9000, 150, 250, 1800],
    "unit_cost": [41000, 1350, 1000, 4300, 6300, 65, 110, 850]
})
n = 6000
df = catalog.iloc[rng.integers(0, len(catalog), n)].reset_index(drop=True).copy()
df.insert(0, "order_id", [f"O{i:05d}" for i in range(1, n+1)])
df.insert(1, "order_date", (
    pd.Timestamp("2025-01-01")
    + pd.to_timedelta(rng.integers(0, 365, n), unit="D")
).strftime("%Y-%m-%d"))
df.insert(2, "customer_id", [f"C{x:04d}" for x in rng.integers(1, 2501, n)])
df.insert(3, "region", rng.choice(
    ["North", "South", "East", "West"], n, p=[.30, .30, .18, .22]))
df["quantity"] = rng.integers(1, 5, n)
df["discount"] = rng.choice(
    [0., .05, .10, .20, .30], n, p=[.30, .20, .25, .15, .10])
df["shipping_cost"] = rng.choice([40, 70, 100, 150], n)

# Inject quality issues before adding exact duplicate rows.
df.loc[:19, "region"] = " " + df.loc[:19, "region"].str.lower() + " "
df.loc[20:29, "region"] = None
raw = pd.concat([df, df.iloc[30:60]], ignore_index=True)
raw.to_csv(ROOT / "data/orders_raw.csv", index=False)
catalog.to_csv(ROOT / "data/products.csv", index=False)

clean = raw.drop_duplicates().copy()
clean["region"] = clean["region"].fillna("Unknown").str.strip().str.title()
clean["order_date"] = pd.to_datetime(
    clean["order_date"], format="%Y-%m-%d", errors="raise")
assert clean["order_id"].is_unique
assert clean[["order_id","customer_id","product","order_date"]].notna().all().all()
assert clean["quantity"].gt(0).all()
assert clean["discount"].between(0,1).all()
clean["gross_sales"] = clean["quantity"] * clean["unit_price"]
clean["revenue"] = (clean["gross_sales"] * (1-clean["discount"])).round(2)
clean["cost"] = clean["quantity"]*clean["unit_cost"] + clean["shipping_cost"]
clean["profit"] = (clean["revenue"] - clean["cost"]).round(2)
clean["month"] = clean["order_date"].dt.strftime("%Y-%m")
clean["discount_band"] = np.select(
    [clean["discount"].eq(0), clean["discount"].le(.1)],
    ["None", "Low (1-10%)"], default="High (>10%)")
clean = clean.sort_values(["order_date","order_id"]).reset_index(drop=True)
clean.to_csv(ROOT / "data/orders_clean.csv", index=False)
quality = {
    "source": "synthetic; seed 42", "currency": "INR",
    "raw_rows": len(raw), "clean_rows": len(clean),
    "exact_duplicates_removed": len(raw)-len(clean),
    "missing_regions_labelled_unknown": int(clean["region"].eq("Unknown").sum())
}
write("reports/data_quality.json", json.dumps(quality, indent=2))

# Relational data loading with an explicit schema.
with sqlite3.connect(ROOT / "sql/retail.db") as conn:
    conn.executescript((ROOT / "sql/schema.sql").read_text())
    catalog.to_sql("products", conn, if_exists="append", index=False)
    export = clean.copy()
    export["order_date"] = export["order_date"].dt.strftime("%Y-%m-%d")
    export.to_sql("orders", conn, if_exists="append", index=False)
    assert not conn.execute("PRAGMA foreign_key_check").fetchall()
    results = {}
    for sql in sorted((ROOT / "sql").glob("[0-9][0-9]_*.sql")):
        result = pd.read_sql_query(sql.read_text(), conn)
        result.to_csv(ROOT / "reports" / (sql.stem+".csv"), index=False)
        results[sql.stem] = result

# Independent Pandas checks: keys, totals, order counts, ratios and repeat segments.
assert len(clean)==6000 and len(raw)-len(clean)==30
assert quality["missing_regions_labelled_unknown"]==10
for key, group in [
    ("01_monthly_growth","month"), ("02_product_profitability","product"),
    ("03_discount_impact","discount_band"), ("05_regional_performance","region")
]:
    expected = clean.groupby(group).agg(
        revenue=("revenue","sum"), profit=("profit","sum"), orders=("order_id","count")
    ).sort_index()
    actual = results[key].set_index(group).sort_index()
    assert actual.index.tolist()==expected.index.tolist()
    np.testing.assert_allclose(
        actual[["revenue","profit","orders"]],
        expected[["revenue","profit","orders"]], atol=.01)
    if "margin_pct" in actual:
        np.testing.assert_allclose(
            actual["margin_pct"], 100*expected["profit"]/expected["revenue"])

monthly = results["01_monthly_growth"]
products = results["02_product_profitability"]
discounts = results["03_discount_impact"]
segments = results["04_customer_segments"]
regions = results["05_regional_performance"]
np.testing.assert_allclose(
    monthly["mom_growth_pct"], monthly["revenue"].pct_change()*100, equal_nan=True)
np.testing.assert_allclose(regions["aov"], regions["revenue"]/regions["orders"])
np.testing.assert_allclose(
    regions["profit_rank"], regions["profit"].rank(method="dense", ascending=False))
loss_counts = clean.assign(loss=clean["profit"].lt(0)).groupby("product")["loss"].sum()
np.testing.assert_array_equal(
    products.set_index("product")["loss_orders"].sort_index(), loss_counts.sort_index())
customer = clean.groupby("customer_id").agg(
    orders=("order_id","count"), revenue=("revenue","sum"), profit=("profit","sum"))
customer["segment"] = np.where(customer["orders"].ge(2), "Repeat", "One-time")
expected_segments = customer.groupby("segment").agg(
    customers=("orders","size"), orders=("orders","sum"),
    revenue=("revenue","sum"), profit=("profit","sum"))
np.testing.assert_allclose(
    segments.set_index("segment").sort_index()[["customers","orders","revenue","profit"]],
    expected_segments.sort_index()[["customers","orders","revenue","profit"]])
loss_expected = clean.groupby("product")["profit"].sum()
loss_expected = loss_expected[loss_expected.lt(0)].sort_index()
loss_actual = results["06_loss_making_products"].set_index("product")["total_profit"].sort_index()
assert loss_expected.index.tolist() == loss_actual.index.tolist()
np.testing.assert_allclose(loss_expected.to_numpy(dtype=float),
    loss_actual.to_numpy(dtype=float),atol=.01)
revenue = float(clean["revenue"].sum())
profit = float(clean["profit"].sum())
repeat = float(customer["orders"].ge(2).mean())

# Excel: five business charts, readable merged headers and a compact phone summary.
with pd.ExcelWriter(ROOT / "excel/retail_analysis.xlsx", engine="xlsxwriter") as writer:
    dash = writer.book.add_worksheet("Dashboard")
    mobile = writer.book.add_worksheet("Mobile Summary")
    clean.to_excel(writer, sheet_name="Orders", index=False)
    for key, result in results.items():
        result.to_excel(writer, sheet_name=key[:31], index=False)
    book = writer.book
    money = book.add_format({"num_format":"#,##0.00"})
    pct = book.add_format({"num_format":"0.0%"})
    title = book.add_format({"font_size":20, "bold":True, "font_color":"#16324F"})
    header = book.add_format({"bold":True,"bg_color":"#16324F","font_color":"white"})
    source = writer.sheets["Orders"]
    source.add_table(0,0,len(clean),len(clean.columns)-1,{
        "name":"Orders", "style":"Table Style Medium 2",
        "columns":[{"header":c} for c in clean.columns]})
    pc = clean.columns.get_loc("profit")
    source.conditional_format(1,pc,len(clean),pc,{"type":"3_color_scale"})
    # Dashboard was created first so it opens before the data sheets.
    dash.activate()
    dash.hide_gridlines(2)
    dash.set_zoom(90)
    dash.set_column("A:A",28)
    dash.set_column("B:H",10)
    for row in range(94):
        dash.set_row(row,20)
    title = book.add_format({"font_size":22,"bold":True,"font_color":"white",
        "bg_color":"#16324F","valign":"vcenter"})
    subtitle = book.add_format({"font_size":11,"font_color":"#52667A","valign":"vcenter"})
    label_fmt = book.add_format({"font_size":12,"bold":True,"font_color":"#16324F",
        "bg_color":"#EDF4F8","valign":"vcenter"})
    note_fmt = book.add_format({"font_size":11,"text_wrap":True,
        "font_color":"#52667A","valign":"vcenter"})
    dash.merge_range("A1:H2","RETAIL ANALYTICS | 2025",title)
    dash.merge_range("A3:H3","Synthetic portfolio data | INR | 6,000 orders",subtitle)
    dash.merge_range("E4:H9",
        "Five business questions\n\nScroll down for all five charts.\n"
        "On a phone, open Mobile Summary for a compact overview.",note_fmt)
    dash.merge_range("A11:H11","FINANCIALS ARE INR; CHART AXES USE MILLIONS WHERE LABELLED",subtitle)
    dash.freeze_panes(3,0)
    dash.set_landscape()
    dash.set_paper(9)
    dash.fit_to_pages(1,0)
    dash.print_area("A1:H93")
    end = len(customer)+1
    metrics = [
        ("Revenue","=SUM(Orders[revenue])",revenue,money),
        ("Profit","=SUM(Orders[profit])",profit,money),
        ("Profit margin","=IFERROR(B5/B4,0)",profit/revenue,pct),
        ("Orders","=COUNTA(Orders[order_id])",len(clean),None),
        ("Average order value","=IFERROR(B4/B7,0)",revenue/len(clean),money),
        ("Repeat customer rate",
         f'=IFERROR(COUNTIF(Customers!B2:B{end},">=2")/COUNTA(Customers!A2:A{end}),0)',
         repeat,pct)]
    for row,(label,formula,value,fmt) in enumerate(metrics,start=3):
        dash.set_row(row,30)
        dash.write(row,0,label,label_fmt)
        value_format = book.add_format({"font_size":18,"bold":True,
            "font_color":"#12788B","bg_color":"#EDF4F8","valign":"vcenter",
            "num_format": "0.0%" if fmt is pct else ("#,##0" if fmt is None else "#,##0.00")})
        dash.merge_range(row,1,row,3,"",value_format)
        dash.write_formula(row,1,formula,value_format,float(value))
    customers = book.add_worksheet("Customers")
    customers.write_row(0,0,["customer_id","orders"],header)
    for row,(cid,item) in enumerate(customer.iterrows(),start=1):
        customers.write(row,0,cid)
        customers.write_formula(row,1,
            f'=COUNTIF(Orders[customer_id],A{row+1})',None,int(item["orders"]))
    region_sheet = book.add_worksheet("Region_Formulas")
    region_sheet.write_row(0,0,["region","revenue","profit","margin"],header)
    for row,item in enumerate(regions.to_dict("records"),start=1):
        erow=row+1
        region_sheet.write(row,0,item["region"])
        for col,field in [(1,"revenue"),(2,"profit")]:
            region_sheet.write_formula(row,col,
                f'=SUMIFS(Orders[{field}],Orders[region],A{erow})',money,item[field])
        region_sheet.write_formula(row,3,f'=IFERROR(C{erow}/B{erow},0)',
            pct,item["margin_pct"]/100)
    chart_specs = [
        ("line","01_monthly_growth",0,2,len(monthly),"A13","1. Monthly revenue","INR millions"),
        ("bar","02_product_profitability",1,4,len(products),"A29","2. Product profit","INR millions"),
        ("column","03_discount_impact",0,5,len(discounts),"A45","3. Margin by discount band","Percent"),
        ("column","04_customer_segments",0,1,len(segments),"A61","4. Repeat and one-time customers","Customers"),
        ("bar","Region_Formulas",0,2,len(regions),"A77","5. Regional profit","INR millions")
    ]
    for ctype,sheet,xcol,ycol,length,position,label,unit in chart_specs:
        chart = book.add_chart({"type":ctype})
        series = {"name":label,"categories":[sheet,1,xcol,length,xcol],
                  "values":[sheet,1,ycol,length,ycol]}
        if ctype=="line":
            series["line"]={"color":"#12788B","width":2.5}
            series["marker"]={"type":"circle","size":5}
        else:
            series["fill"]={"color":"#12788B"}
            series["border"]={"none":True}
        chart.add_series(series)
        chart.set_title({"name":label,"name_font":{"size":15,"color":"#16324F"}})
        chart.set_legend({"none":True})
        chart.set_size({"width":695,"height":290})
        chart.set_chartarea({"border":{"none":True},"fill":{"color":"#FFFFFF"}})
        number_axis={"name":unit,"num_font":{"size":10},
                     "num_format": '0.0,,"M"' if unit=="INR millions" else "0.0"}
        if ctype=="bar":
            chart.set_x_axis(number_axis)
            chart.set_y_axis({"num_font":{"size":11}})
        else:
            chart.set_y_axis(number_axis)
            chart.set_x_axis({"num_font":{"size":10},"label_position":"low"})
        dash.insert_chart(position,chart)

    # Mobile Summary was created second.
    mobile.hide_gridlines(2)
    mobile.set_column("A:A",25)
    mobile.set_column("B:B",22)
    mobile.set_zoom(100)
    mobile.set_row(0,35)
    mobile.merge_range("A1:B1","RETAIL | 2025",title)
    mobile.set_row(1,26)
    mobile.merge_range("A2:B2","Synthetic data | INR",subtitle)
    for row,(label,formula,value,fmt) in enumerate(metrics,start=3):
        mobile.set_row(row,32)
        mobile.write(row,0,label,label_fmt)
        mobile.write_formula(row,1,f"=Dashboard!B{row+1}",fmt,float(value))
    notes=[
        ("1. Sales trend",f"Peak month: {monthly.loc[monthly['revenue'].idxmax(),'month']}."),
        ("2. Product profitability",f"Lowest total profit: {products.iloc[0]['product']}. "
         f"Negative-profit products: {len(results['06_loss_making_products'])}."),
        ("3. Discounts",f"Lowest-margin band: {discounts.iloc[0]['discount_band']} "
         f"({discounts.iloc[0]['margin_pct']:.2f}%)."),
        ("4. Repeat buying",f"{int(customer['orders'].ge(2).sum()):,} of "
         f"{len(customer):,} customers bought at least twice."),
        ("5. Regions",f"Highest total profit: {regions.iloc[0]['region']}. "
         "Compare margin and volume before taking action."),
        ("Scope","Contribution profit excludes overhead, tax, returns and ads. "
         "These synthetic results are not real business outcomes.")
    ]
    for i,(label,note) in enumerate(notes):
        row=11+i*3
        mobile.set_row(row,24)
        mobile.merge_range(row,0,row,1,label,label_fmt)
        mobile.set_row(row+1,50)
        mobile.merge_range(row+1,0,row+1,1,note,note_fmt)
    mobile.print_area("A1:B29")
    mobile.fit_to_pages(1,1)
    for sheet in writer.sheets.values():
        if sheet.name not in {"Dashboard","Mobile Summary"}:
            sheet.freeze_panes(1,0)
            sheet.set_column(0,20,19)


# Validate workbook structures and stored formula results; not Excel rendering.
wb = openpyxl.load_workbook(ROOT / "excel/retail_analysis.xlsx",data_only=True)
np.testing.assert_allclose(
    [wb["Dashboard"][f"B{i}"].value for i in range(4,10)],
    [revenue,profit,profit/revenue,len(clean),revenue/len(clean),repeat])
for row,item in enumerate(regions.to_dict("records"),start=2):
    np.testing.assert_allclose(
        [wb["Region_Formulas"].cell(row,c).value for c in [2,3,4]],
        [item["revenue"],item["profit"],item["margin_pct"]/100])
assert wb["Orders"].max_row==6001
np.testing.assert_allclose(
    [wb["Mobile Summary"][f"B{i}"].value for i in range(4,10)],
    [revenue,profit,profit/revenue,len(clean),revenue/len(clean),repeat])
wb.close()
wf = openpyxl.load_workbook(ROOT / "excel/retail_analysis.xlsx",data_only=False)
assert wf["Dashboard"]["B4"].data_type=="f"
assert len(wf["Dashboard"]._charts)==5
assert "A1:H2" in {str(r) for r in wf["Dashboard"].merged_cells.ranges}
assert wf["Dashboard"].row_dimensions[1].height >= 20
assert wf.sheetnames[:2] == ["Dashboard","Mobile Summary"]
assert wf["Dashboard"].sheet_view.showGridLines is False
assert wf["Mobile Summary"]["B4"].data_type=="f"
wf.close()

peak=monthly.loc[monthly["revenue"].idxmax()]
worst=products.iloc[0]
low=discounts.iloc[0]
best=regions.iloc[0]
loss=results["06_loss_making_products"]
write("reports/executive_summary.md", f"""
# Executive summary

Synthetic educational dataset; computed results, not actual company performance.

| KPI | Value |
|---|---:|
| Orders | {len(clean):,} |
| Observed customers | {len(customer):,} |
| Revenue | INR {revenue:,.2f} |
| Illustrative profit | INR {profit:,.2f} |
| Weighted margin | {profit/revenue:.2%} |
| Average order value | INR {revenue/len(clean):,.2f} |
| Repeat customer rate | {repeat:.2%} |

## Five answers

1. **Sales trend:** peak revenue is {peak['month']} at INR {peak['revenue']:,.2f}.
   See 01_monthly_growth.csv for monthly growth. First-month growth is undefined.
   Review staffing around peaks; one year cannot prove recurring seasonality.
2. **Products:** {worst['product']} has the lowest total profit at
   INR {worst['profit']:,.2f}. {len(loss)} products have negative total profit.
   Inspect loss orders and unit economics before changing prices or assortment.
   Lowest profit is not the same as negative profit.
3. **Discounts:** {low['discount_band']} has the lowest weighted margin,
   {low['margin_pct']:.2f}%. Test discount limits within product categories.
   Revenue mechanically falls with discount at fixed quantity; this comparison
   does not estimate demand response or causal campaign benefit.
4. **Customers:** {int(customer['orders'].ge(2).sum()):,} of
   {len(customer):,} observed customers ordered at least twice ({repeat:.2%}).
   Propose a randomized second-purchase campaign; this is not cohort retention.
5. **Regions:** {best['region']} ranks first on total profit
   (INR {best['profit']:,.2f}). Compare margin and AOV alongside volume before
   reallocating resources. Unknown regions remain in the totals.

## Scope and limitations

Calendar year 2025; INR; one single-product order per row.
Profit subtracts product and shipping costs only. No returns, tax, marketing or
fixed overhead. Synthetic patterns cannot justify real market decisions.
Recommendations are untested hypotheses. No percentage improvement is claimed.
""")

# Execute the portable notebook so GitHub shows actual analysis outputs.
nb = nbformat.v4.new_notebook()
nb.cells = [
    nbformat.v4.new_markdown_cell(
        "# Retail analytics exploration\nSynthetic data. Read the data dictionary first."),
    nbformat.v4.new_code_cell(
        "from pathlib import Path\nimport pandas as pd\n"
        "root = Path.cwd()\nif not (root/'data').exists(): root=root.parent\n"
        "df = pd.read_csv(root/'data/orders_clean.csv',parse_dates=['order_date'])\ndf.head()"),
    nbformat.v4.new_code_cell("df.groupby('month')[['revenue','profit']].sum()"),
    nbformat.v4.new_code_cell(
        "p=df.groupby('product')[['revenue','profit']].sum()\n"
        "p['margin_pct']=100*p['profit']/p['revenue']\np.sort_values('profit')"),
    nbformat.v4.new_code_cell(
        "c=df.groupby('customer_id')['order_id'].nunique()\n"
        "print('Repeat customer rate:',c.ge(2).mean())"),
    nbformat.v4.new_markdown_cell(
        "Exercises: compare the discount and region outputs below with their SQL CSV files. "
        "Explain why discount comparisons are not causal evidence.")
]
nb.cells.extend([
    nbformat.v4.new_markdown_cell("## Discount analysis"),
    nbformat.v4.new_code_cell(
        "d=df.groupby('discount_band')[['revenue','profit']].sum()\n"
        "d['margin_pct']=100*d['profit']/d['revenue']\nd"),
    nbformat.v4.new_markdown_cell("## Regional analysis"),
    nbformat.v4.new_code_cell(
        "r=df.groupby('region').agg(revenue=('revenue','sum'),"
        "profit=('profit','sum'),orders=('order_id','count'))\n"
        "r['margin_pct']=100*r['profit']/r['revenue']\n"
        "r['aov']=r['revenue']/r['orders']\nr.sort_values('profit',ascending=False)")
])
nb.metadata["kernelspec"]={"display_name":"Python 3","language":"python","name":"python3"}
NotebookClient(nb,timeout=120,kernel_name="python3",
    resources={"metadata":{"path":str(ROOT)}}).execute()
nbformat.write(nb, ROOT / "python/analysis.ipynb")

write("reports/validation.json",json.dumps({
    "status":"passed",
    "python":platform.python_version(),"pandas":pd.__version__,
    "numpy":np.__version__,"plotly":plotly.__version__,
    "xlsxwriter":xlsxwriter.__version__,"sqlite":sqlite3.sqlite_version,
    "checks":[
        "6000 unique orders; 30 duplicates removed; 10 unknown regions",
        "required fields, positive quantity, discount bounds, foreign keys",
        "SQL/Pandas grouped financial totals and order counts",
        "weighted margins, MoM growth, AOV, dense rank, loss order counts",
        "SQL/Pandas customer segments: counts, orders, revenue and profit",
        "Excel KPI, mobile and regional cached values; five charts; merged title and sheet order",
        "Notebook code cells executed successfully"
    ],
    "not_validated":["native Power BI PBIX", "desktop Excel visual rendering",
                     "browser dashboard visual rendering"]
},indent=2))
from portfolio import build_portfolio
build_portfolio(ROOT, clean, results)

archive=ROOT/"dist/retail-analytics.zip"
with zipfile.ZipFile(archive,"w",zipfile.ZIP_DEFLATED) as z:
    for path in sorted(ROOT.rglob("*")):
        rel=path.relative_to(ROOT)
        if path.is_file() and not any(
            p in {".git",".venv","__pycache__","dist"} for p in rel.parts
        ):
            z.write(path,arcname=str(rel))
print(json.dumps({"status":"passed","orders":len(clean),"revenue":revenue,
                  "profit":profit,"repeat_rate":repeat,"zip":str(archive)},indent=2))
